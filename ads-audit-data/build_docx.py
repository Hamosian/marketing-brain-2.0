from docx import Document
from docx.shared import Pt, RGBColor, Inches, Twips
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from pathlib import Path

# Resolve asset/output paths relative to this script, not the process CWD
BASE_DIR = Path(__file__).resolve().parent

# ---- Brand constants ----
PURPLE="7C5CFF"; NEAR_BLACK="0F0F14"; WHITE="FFFFFF"; DARK_GRAY="1C1C24"
MID_GRAY="2A2A35"; LIGHT_GRAY="E6E6EB"; PURPLE_LIGHT="EDE8FF"; PURPLE_MIST="F7F5FF"
FONT="Inter"

def C(h): return RGBColor.from_string(h)

doc=Document()
# base style
st=doc.styles['Normal']; st.font.name=FONT; st.font.size=Pt(10); st.font.color.rgb=C(NEAR_BLACK)
st.paragraph_format.space_after=Pt(6); st.paragraph_format.line_spacing=1.12

# margins
for s in doc.sections:
    s.top_margin=Inches(0.7); s.bottom_margin=Inches(0.7)
    s.left_margin=Inches(0.7); s.right_margin=Inches(0.7)

def shade(cell,hexc):
    tcPr=cell._tc.get_or_add_tcPr()
    sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),hexc)
    tcPr.append(sh)

def cell_margins(cell,t=60,b=60,l=100,r=100):
    tcPr=cell._tc.get_or_add_tcPr()
    m=OxmlElement('w:tcMar')
    for tag,val in (('w:top',t),('w:bottom',b),('w:start',l),('w:end',r)):
        e=OxmlElement(tag); e.set(qn('w:w'),str(val)); e.set(qn('w:type'),'dxa'); m.append(e)
    tcPr.append(m)

def set_borders(tbl):
    tblPr=tbl._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement(f'w:{edge}'); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'8')
        e.set(qn('w:space'),'0'); e.set(qn('w:color'),LIGHT_GRAY); borders.append(e)
    tblPr.append(borders)

def runc(p,text,color=NEAR_BLACK,bold=False,size=10,italic=False):
    r=p.add_run(text); r.font.name=FONT; r.font.size=Pt(size); r.font.bold=bold; r.font.italic=italic
    r.font.color.rgb=C(color); return r

def h1(text):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(14); p.paragraph_format.space_after=Pt(2)
    runc(p,text,NEAR_BLACK,True,16)
    # purple accent bar
    bar=doc.add_paragraph(); bar.paragraph_format.space_after=Pt(8); bar.paragraph_format.space_before=Pt(0)
    pPr=bar._p.get_or_add_pPr(); pbdr=OxmlElement('w:pBdr')
    bottom=OxmlElement('w:bottom'); bottom.set(qn('w:val'),'single'); bottom.set(qn('w:sz'),'18')
    bottom.set(qn('w:space'),'1'); bottom.set(qn('w:color'),PURPLE); pbdr.append(bottom); pPr.append(pbdr)
    r=bar.add_run(); r.font.size=Pt(2)

def h2(text):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(10); p.paragraph_format.space_after=Pt(3)
    runc(p,text,PURPLE,True,12.5)

def h3(text):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(6); p.paragraph_format.space_after=Pt(2)
    runc(p,text,DARK_GRAY,True,11)

def body(text,size=10):
    p=doc.add_paragraph(); runc(p,text,NEAR_BLACK,False,size); return p

def label_body(label,text,size=10):
    p=doc.add_paragraph(); runc(p,label,PURPLE,True,size); runc(p,text,NEAR_BLACK,False,size); return p

def bullet(text,size=10,bold_lead=None):
    p=doc.add_paragraph(style='List Bullet')
    if bold_lead: runc(p,bold_lead,PURPLE,True,size)
    runc(p,text,NEAR_BLACK,False,size); 
    p.paragraph_format.space_after=Pt(3)
    return p

def numbered(text,size=10,bold_lead=None):
    p=doc.add_paragraph(style='List Number')
    if bold_lead: runc(p,bold_lead,PURPLE,True,size)
    runc(p,text,NEAR_BLACK,False,size)
    p.paragraph_format.space_after=Pt(3)
    return p

def table(headers,rows,widths=None,fs=8.5,header_fs=8.5):
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    set_borders(t)
    hc=t.rows[0].cells
    for i,htext in enumerate(headers):
        shade(hc[i],NEAR_BLACK); cell_margins(hc[i])
        para=hc[i].paragraphs[0]; para.paragraph_format.space_after=Pt(0)
        runc(para,htext,WHITE,True,header_fs)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for ci,val in enumerate(row):
            if ri%2==1: shade(cells[ci],PURPLE_LIGHT)
            cell_margins(cells[ci])
            para=cells[ci].paragraphs[0]; para.paragraph_format.space_after=Pt(0)
            runc(para,str(val),NEAR_BLACK,False,fs)
    if widths:
        for i,w in enumerate(widths):
            for r in t.rows:
                r.cells[i].width=Inches(w)
    return t

# ---- Header / footer ----
sec=doc.sections[0]
hdr=sec.header; hp=hdr.paragraphs[0]; hp.alignment=WD_ALIGN_PARAGRAPH.RIGHT
runc(hp,"Riverside.com",PURPLE,True,9); runc(hp,"  paid ads audit",MID_GRAY,False,9)
ftr=sec.footer; fp=ftr.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
runc(fp,"Confidential  |  Marketing OS agent  |  2026-06-24  |  page ",MID_GRAY,False,8)
# page number field
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); 
r=OxmlElement('w:r'); rpr=OxmlElement('w:rPr'); fp._p.append(fld); fld.append(r)

# ---- Title block ----
logo_p=doc.add_paragraph(); logo_p.alignment=WD_ALIGN_PARAGRAPH.LEFT
logo_p.add_run().add_picture(str(BASE_DIR/"riverside_logo.png"),width=Inches(1.9))
tp=doc.add_paragraph(); tp.paragraph_format.space_before=Pt(8); tp.paragraph_format.space_after=Pt(2)
runc(tp,"Paid ads audit: Google and Meta",NEAR_BLACK,True,22)
sub=doc.add_paragraph(); sub.paragraph_format.space_after=Pt(1)
runc(sub,"30-day performance refresh",PURPLE,True,12)
meta=doc.add_paragraph(); meta.paragraph_format.space_after=Pt(2)
runc(meta,"Date 2026-06-24   |   Window last 30 days (2026-05-25 to 2026-06-24)   |   Source Windsor.ai live pull   |   Prepared by the Marketing OS agent",MID_GRAY,False,9)
# accent bar under title
bar=doc.add_paragraph(); pPr=bar._p.get_or_add_pPr(); pbdr=OxmlElement('w:pBdr')
bottom=OxmlElement('w:bottom'); bottom.set(qn('w:val'),'single'); bottom.set(qn('w:sz'),'24'); bottom.set(qn('w:space'),'1'); bottom.set(qn('w:color'),PURPLE); pbdr.append(bottom); pPr.append(pbdr)
bar.add_run().font.size=Pt(2)

label_body("Scope. ","Focused 30-day Google and Meta refresh. It complements, and does not replace, the full 90-day all-platform audit (2026-06-14) that also covers LinkedIn and Microsoft/Bing.",9)
label_body("Currency. ","Figures are in account-reported currency and are directional until FX is confirmed (some Riverside accounts may bill in ILS). Per-platform budget shares and all campaign-versus-campaign comparisons hold regardless.",9)
label_body("Method. ","Windsor exposes performance data, not account configuration. Scores reflect only the data-observable subset (Google 15 of 80 checks, Meta 12 of 50). Configuration checks are flagged for platform-UI verification, not guessed.",9)

# ---- Executive summary ----
h1("Executive summary")
table(
 ["Scope","Score","Grade","Spend share (30d)"],
 [["Google Ads","65.0 / 100","C  Needs improvement","98.1%"],
  ["Meta Ads","53.5 / 100","D  Poor (signal layer F)","1.9%"],
  ["Aggregate (budget-weighted)","64.8 / 100","C","100%"]],
 widths=[2.3,1.3,2.4,1.3], fs=9.5, header_fs=9.5)
body("")
body("Meta is under 2 percent of paid spend, so the blended score is effectively the Google score. Both platforms are reported separately so Meta's problems are not masked. Total 30-day spend across both platforms is 1,210,992 in account currency.",9.5)

h2("What changed since the 2026-06-14 audit")
body("The most important finding of this refresh: the prior audit's top issue is unaddressed and deteriorating.",9.5)
table(
 ["Item","2026-06-14 (90d)","2026-06-24 (30d)","Direction"],
 [["US_B2C_Broad CPA","592","932","Worse by 57%"],
  ["US_B2C_Broad spend pace","~150k / 30d","172.5k / 30d","Up"],
  ["US_B2C_Broad bidding","Max Conv, no tCPA cap","Still no tCPA cap","Unchanged"],
  ["Meta purchase value and Lead event","Both null","Both still null","Unchanged"],
  ["Meta worst retargeting frequency","24.5","16.7","Better but still failing"]],
 widths=[1.9,2.0,1.9,1.3], fs=9, header_fs=9)

h2("Top critical issues")
numbered("172,512 in 30 days at a 932 cost per conversion versus 20 to 70 on strong campaigns. Search impression share 22 percent, so it loses most auctions while still spending heavily. Flagged 10 days ago, now worse. (Google)",9.5,bold_lead="US_B2C_Broad_Desktop is the largest line item and the worst performer. ")
numbered("8 enabled campaigns above 200 CPA hold 457,686 in 30 days (39 percent of Google spend) while comparable campaigns convert at 20 to 70. (Google)",9.5,bold_lead="The high-CPA cluster. ")
numbered("102,445 over 30 days went to search terms with over 10 spend and zero conversions (video editor, webinar, teleprompter online, opus clip). That is 17.3 percent of meaningful search-term spend. (Google)",9.5,bold_lead="Search-term waste. ")
numbered("Purchase value is 0 and the website Lead event is 0 on every Meta campaign, so Meta optimizes on a custom event with no revenue or lead signal. (Meta and cross-platform)",9.5,bold_lead="Meta cannot optimize for value or leads. ")
numbered("Campaigns run Maximize Conversions (volume), not value, so the system optimizes toward signups regardless of downstream worth. (Google)",9.5,bold_lead="Conversion value missing on ~99 percent of non-brand Google spend. ")

h2("Top quick wins (high impact, under 15 minutes)")
table(
 ["Fix","Platform","Severity","Time"],
 [["Add a Target CPA cap to US_B2C_Broad_Desktop (start 60 to 90), or pause it","Google","Critical","10 min"],
  ["Add negative keywords for the top zero-conversion search terms","Google","Critical","10 min"],
  ["Add tCPA caps to the other uncapped high-CPA campaigns","Google","Critical","15 min"],
  ["Pause UK_PMax_Podcasting and US_PMax_Podcasting_Clean","Google","High","2 min"],
  ["Switch Target Spend competitor campaigns to Maximize Conversions","Google","High","5 min each"],
  ["Add a frequency cap on Meta Remarket > Event_SignedUp (16.7)","Meta","High","10 min"],
  ["Confirm Enhanced Conversions is enabled and verified","Google","High","5 min"]],
 widths=[4.2,1.0,1.1,1.0], fs=9, header_fs=9)

# ---- Google ----
doc.add_page_break()
h1("Google Ads")
label_body("Score 65.0 / 100, Grade C. ","Scored on 15 of 80 observable checks. Strong account hygiene (naming, brand separation, smart bidding on most campaigns, clean keyword-level waste) is dragged down by search-term waste, missing conversion values, and a small cluster of very expensive campaigns.",9.5)
h3("Category breakdown (observable subset)")
table(
 ["Category","Weight","Verdict"],
 [["Conversion tracking","25%","Mixed. Conversions fire everywhere (pass). Conversion value absent on ~99% of non-brand spend (fail). Enhanced Conversions, Consent Mode, server-side, attribution: needs UI."],
  ["Wasted spend and negatives","20%","Weak. 17.3% of search-term spend on zero-conversion terms (fail). Keyword-level waste is clean (pass). Broad match paired with smart bidding (pass)."],
  ["Account structure","15%","Good. Clear naming (pass). Brand and non-brand separated via a dedicated Brand account (pass). PMax present (pass). Fragmentation and skewed allocation (warning)."],
  ["Keywords and Quality Score","15%","Not scored. Windsor returns a summed Quality Score. Verify QS and component ratings in the UI."],
  ["Ads and assets","15%","Partial. Search CTRs strong (pass). RSA strength, PMax asset density and video: needs UI."],
  ["Settings and targeting","10%","Mixed. Target Spend on some competitor and retargeting campaigns (warning). Brand on Manual CPC with high volume, defensible but flagged. Several high-spend campaigns budget-limited (warning)."]],
 widths=[1.9,0.7,4.5], fs=8.5, header_fs=9)

h3("The high-CPA cluster (kill or fix list)")
body("8 enabled campaigns, 457,686 in 30 days (39 percent of Google spend), all above 200 CPA while comparable campaigns convert at 20 to 70.",9)
table(
 ["Campaign","Spend 30d","CPA","Search IS","Action"],
 [["US_B2C_Broad_Desktop","172,512","932","22%","Add tCPA or pause and rebuild. ~13x peer CPA. Worse since last audit."],
  ["US_PMax_DescriptEditing_Desktop","69,282","409","15%","Budget-starved and expensive. Restructure or pause."],
  ["US_Generic_Webinar_Desktop","58,827","222","56%","Has some value. Tighten targeting and negatives."],
  ["UK_PMax_Podcasting_Desktop","51,987","781","19%","Pause. Very low volume at very high CPA."],
  ["US_Generic_Podcast_Recording_Desktop","38,027","207","82%","High IS, not budget-limited. The CPA is the issue."],
  ["CAUKAU_B2C_Alpha_Desktop","36,733","224","72%","Review targeting and bids."],
  ["TopGEOs_B2C_Alpha_Mobile","16,045","452","94%","Near-full IS at high CPA. Check mobile landing page."],
  ["US_PMax_Podcasting_Desktop_Clean","14,274","510","27%","Pause. Lowest volume, high CPA."]],
 widths=[2.5,0.9,0.6,0.7,2.4], fs=8, header_fs=8.5)
body("Caveat: CPA comparison assumes a consistent conversion definition across campaigns. B2C Broad and Alpha campaigns may optimize toward a different conversion event, which would partly explain the gap. Even allowing for that, 932 and 781 are extreme.",8.5)

h3("Top wasted search terms (over 10 spend, 0 conversions, 30 days)")
table(
 ["Spend","Clicks","Campaign","Search term"],
 [["1,085","325","US_B2C_Broad_Desktop","video editor"],
  ["847","65","US_Generic_Webinar_Mobile","webinar"],
  ["776","47","US_Generic_Webinar_Mobile","webinarjam"],
  ["521","163","US_B2C_Broad_Desktop","teleprompter online"],
  ["478","25","TopGEOs_B2C_Alpha_Mobile","best podcast hosting platform"],
  ["472","30","US_B2C_Broad_Desktop","podcast studio"],
  ["452","73","US_B2C_Broad_Desktop","opus clip"],
  ["443","21","US_B2C_Broad_Desktop","apple podcast connect"]],
 widths=[0.9,0.8,3.0,2.4], fs=8.5, header_fs=8.5)
body("Higher-confidence subset: 363 terms with over 50 spend and 0 conversions = 39,020 in 30 days. Many leaking terms are off-intent, which points to thin negative-keyword coverage on the broad and DSA campaigns.",8.5)

# ---- Meta ----
doc.add_page_break()
h1("Meta Ads")
label_body("Score 53.5 / 100, Grade D. ","Scored on 12 of 50 observable checks. The foundational signal layer is F-grade in isolation (purchase value and Lead event both dead, value-based bidding impossible), consistent with the 2026-06-14 audit that scored Meta 7/100. The higher number here reflects a narrower observable check set plus genuinely strong creative on the podcasting campaign. It should not be read as Meta is fine.",9.5)
h3("Category breakdown (observable subset)")
table(
 ["Category","Weight","Verdict"],
 [["Pixel and CAPI health","30%","Pixel firing (pass). Account optimizes on a custom pixel event while the standard Lead event returns 0 and purchase value returns 0 (fail). CAPI, EMQ, dedup, domain, AEM: needs Events Manager."],
  ["Creative (diversity and fatigue)","30%","Weak. Only 2 formats account-wide, no carousel, several single-format ad sets (fail). 9:16 vertical present on podcasting (pass). App and Brand ad sets carry only 3 creatives (warning)."],
  ["Account structure","20%","Acceptable. 5 campaigns, above the 1 to 3 ideal but each a distinct objective (warning). CBO on ad sets above 100/day (pass). Ad sets above 10/day (pass)."],
  ["Audience and targeting","20%","Not scored. Overlap, freshness, lookalike quality, exclusions, first-party uploads: UI only. Verify purchasers are excluded from prospecting."]],
 widths=[1.9,0.7,4.5], fs=8.5, header_fs=9)

h3("Spend and signal by campaign (30 days)")
table(
 ["Campaign","Objective","Spend","Custom ev.","Regs","Purch.","Purch. value","Lead"],
 [["CBO_Podcasting_US","Sales","7,195","247","65","9","0","0"],
  ["Remarket_Visitors_WW","Sales","6,289","9,249","449","648","0","0"],
  ["Newsletter_CBO","Sales","4,488","315","16","0","0","0"],
  ["App_iOS_US","App installs","2,402","8","93","0","0","0"],
  ["Brand_Love-Campaign","Engagement","2,035","4","4","0","0","0"]],
 widths=[1.8,1.1,0.8,0.85,0.55,0.6,0.95,0.5], fs=8, header_fs=8)
body("Purchase value is 0 and Lead is 0 on every row. The remarketing campaign reports 648 purchases with no value, so ROAS cannot be measured anywhere on Meta.",8.5)

h3("Frequency and creative detail")
table(
 ["Campaign > ad set","Freq (30d)","Read"],
 [["Remarket > Event_SignedUp","16.7","Severe overexposure. Cap to 8 to 12. Improved from 24.5 but still failing."],
  ["Remarket > Website-Visitors","7.2","Acceptable for retargeting, watch it."],
  ["Newsletter > US_NewsletterSignup","4.9","Prospecting frequency over the 3.0 target (warning)."],
  ["CBO_Podcasting > YouTubeRepurpose","2.3","Healthy. 8 vertical video concepts, 9 to 15% hook CTR. Best creative in the account."],
  ["App > AppInstalls","1.5","Healthy."],
  ["Brand_Love > Feeds-Reels","1.1","Healthy. CTR 0.25% (engagement objective)."]],
 widths=[2.6,0.9,3.6], fs=8.5, header_fs=8.5)

# ---- Cross-platform ----
h1("Cross-platform and recommendations")
h3("Cross-platform checks")
table(
 ["Check","Verdict","Detail"],
 [["Privacy and signal infrastructure","Fail (critical)","Meta purchase value and Lead event dead. Google conversion value absent on ~99% of non-brand spend. Consent Mode v2 unverified. Neither platform can optimize toward revenue."],
  ["Creative diversity","Warning (high)","Meta runs only 2 formats with no carousel. Google PMax asset density unverified."],
  ["Refresh cadence","Not scored","Creative launch dates not exposed via Windsor. Some Meta concepts appear long-running."]],
 widths=[1.9,1.4,3.8], fs=8.5, header_fs=9)

h3("Strategic recommendations")
numbered("Pass conversion value into Google and fix the Meta Lead event and purchase value. Until then, every bidding decision on both platforms is blind to revenue.",9.5,bold_lead="Fix measurement before scaling anything. ")
numbered("Adding tCPA to, or pausing, the 8 campaigns above 200 CPA addresses 39 percent of Google spend. Start with US_B2C_Broad (932) and UK_PMax_Podcasting (781). This was the prior audit's top recommendation and is still open.",9.5,bold_lead="Cap or cut the high-CPA cluster. ")
numbered("Build themed negative-keyword lists (generic video editing, webinar tools, hardware, off-intent competitors) and apply at account level. Target the broad and DSA campaigns first.",9.5,bold_lead="Plug the search-term leak. ")
numbered("This data audit covered the observable subset. Quality Score, Enhanced Conversions, Consent Mode, CAPI and EMQ, negative lists, assets, audience exclusions, and landing pages still need a UI pass.",9.5,bold_lead="Run a configuration audit in the platform UIs. ")

p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(10)
runc(p,"Owners: Raz Navon (Head of Paid Acquisition), Dor Druker (Growth Channels Lead).  Slack #marketing-growth-ppc-team.  Companion files: ADS-AUDIT-REPORT-2026-06-24.md, ADS-ACTION-PLAN-2026-06-24.md, ADS-QUICK-WINS-2026-06-24.md.",MID_GRAY,False,8.5)

out=BASE_DIR.parent/"Riverside-Paid-Ads-Audit-2026-06-24.docx"
doc.save(str(out))
print("saved",out)
