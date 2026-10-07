"""Riverside-branded .docx builder.

Why .docx and not HTML: Google Drive's HTML importer keeps character-level run
styling (colour, size, weight) and table cell shading, and throws away
block-level box styling - margins, padding, borders on div/p, and any custom
<hr>. So an HTML document styled with CSS arrives in Google Docs with Riverside
colours and none of Riverside's layout: no paragraph spacing, a grey hairline
where the purple accent bar should be, and callout boxes flattened into
highlighted text.

.docx converts through Word's document model instead, which has real paragraph
spacing, table borders, cell shading, keep-with-next and repeating header rows.
Everything below survives the conversion to a Google Doc.

Palette and rules per .claude/skills/riverside-brand-guidelines (docx section).

Getting the result into Google Docs, two routes:

  1. Copy the .docx into the user's synced Drive folder
     (~/Library/CloudStorage/GoogleDrive-<email>/My Drive), wait for it to sync,
     open it at https://docs.google.com/document/d/<fileId>/edit and use
     File > Save as Google Docs. No payload passes through the model context.
  2. Upload with the Drive connector's create_file using base64Content and the
     .docx content type; Drive converts on upload. Run slim() first - python-docx
     ships ~370 unused style definitions and a legacy stylesWithEffects part, and
     dropping them takes a typical document from ~45KB to ~19KB.

Written 2026-09-16 after a project brief and one-pager were published as styled
HTML and arrived with Riverside's colours and none of its layout.
"""

import copy

from docx import Document
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Inches, Emu, Twips

PURPLE = RGBColor(0x7C, 0x5C, 0xFF)
NEAR_BLACK = RGBColor(0x0F, 0x0F, 0x14)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK_GRAY = RGBColor(0x1C, 0x1C, 0x24)
MID_GRAY = RGBColor(0x2A, 0x2A, 0x35)

HEX_PURPLE = "7C5CFF"
HEX_NEAR_BLACK = "0F0F14"
HEX_LIGHT_GRAY = "E6E6EB"
HEX_PURPLE_LIGHT = "EDE8FF"
HEX_PURPLE_MIST = "F7F5FF"

FONT = "Arial"  # Instrument Sans is not available inside Google Docs.

# CT_TblPr is a strict sequence. python-docx knows it but deletes its copy
# (_tag_seq) at import, so appending a child lands it after w:tblLook and Word
# can report the file as unreadable content.
TBLPR_ORDER = (
    "tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
    "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders",
    "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription",
    "tblPrChange",
)


# ---------------------------------------------------------------- low level

def _shade(el, hex_fill):
    """Cell or paragraph shading. ShadingType CLEAR, never SOLID."""
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    el.append(shd)


def _cell_shade(cell, hex_fill):
    _shade(cell._tc.get_or_add_tcPr(), hex_fill)


def _cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    mar = OxmlElement("w:tcMar")
    for tag, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        node = OxmlElement(f"w:{tag}")
        node.set(qn("w:w"), str(val))
        node.set(qn("w:type"), "dxa")
        mar.append(node)
    tcPr.append(mar)


def _tblPr_insert(tblPr, el):
    """Insert a tblPr child at its schema position, not at the end."""
    tag = el.tag.split("}")[1]
    after = TBLPR_ORDER[TBLPR_ORDER.index(tag) + 1:]
    tblPr.insert_element_before(el, *(f"w:{t}" for t in after))


def _table_borders(table, hex_color=HEX_LIGHT_GRAY, sz=8):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), str(sz))
        e.set(qn("w:space"), "0")
        e.set(qn("w:color"), hex_color)
        borders.append(e)
    _tblPr_insert(tblPr, borders)


def _grid_widths(table, widths):
    """Write column widths (inches) into w:tblGrid and w:tblW.

    Per-cell tcW is not enough: Pages and Google Docs size columns from tblGrid,
    which python-docx leaves at an equal split, so a check column came out as
    wide as the text beside it. tblW is edited in place because python-docx has
    already placed it, first in tblPr's sequence.
    """
    tbl = table._tbl
    twips = [int(round(w * 1440)) for w in widths]
    for col, w in zip(tbl.tblGrid.findall(qn("w:gridCol")), twips):
        col.set(qn("w:w"), str(w))
    tblW = tbl.tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        _tblPr_insert(tbl.tblPr, tblW)
    tblW.set(qn("w:w"), str(sum(twips)))
    tblW.set(qn("w:type"), "dxa")


def _row_keep_together(row, header=False):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))
    if header:
        trPr.append(OxmlElement("w:tblHeader"))


def _left_bar(cell, hex_color=HEX_PURPLE, sz=24):
    """A thick coloured left border on one cell: the callout bar."""
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), str(sz))
    left.set(qn("w:space"), "0")
    left.set(qn("w:color"), hex_color)
    borders.append(left)
    for edge in ("top", "bottom", "right"):
        e = OxmlElement(f"w:{edge}")
        e.set(qn("w:val"), "nil")
        borders.append(e)
    tcPr.append(borders)


def add_hyperlink(paragraph, url, text, size=10):
    part = paragraph.part
    r_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = OxmlElement("w:r")
    rPr = OxmlElement("w:rPr")
    for tag, val in (("w:color", HEX_PURPLE),):
        el = OxmlElement(tag)
        el.set(qn("w:val"), val)
        rPr.append(el)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rPr.append(u)
    rf = OxmlElement("w:rFonts")
    rf.set(qn("w:ascii"), FONT)
    rf.set(qn("w:hAnsi"), FONT)
    rPr.append(rf)
    sz_el = OxmlElement("w:sz")
    sz_el.set(qn("w:val"), str(int(size * 2)))
    rPr.append(sz_el)
    run.append(rPr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    link.append(run)
    paragraph._p.append(link)


# --------------------------------------------------------------- high level

def new_doc(margin_in=1.0):
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(10)
    normal.font.color.rgb = NEAR_BLACK
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(margin_in)
        s.left_margin = s.right_margin = Inches(margin_in)
    _brand_heading_styles(doc)
    return doc


HEADING_LOOK = {  # level: (size pt, colour, space before, space after)
    1: (18, NEAR_BLACK, 18, 4),
    2: (14, PURPLE, 18, 6),
    3: (11, DARK_GRAY, 12, 4),
}


def _brand_heading_styles(doc):
    """Put the brand on the Heading 1-3 *styles*, not only on the runs.

    h1/h2/h3 use real Heading styles so Word shows a navigation pane and Google Docs
    builds a document outline, which a 15-page doc needs. python-docx's default
    template gives those styles theme fonts (asciiTheme etc.) and a blue theme
    colour; a theme attribute can win over a run's plain font, so they are stripped
    and replaced with the brand font and colour."""
    for level, (_, color, _, _) in HEADING_LOOK.items():
        st = doc.styles[f"Heading {level}"]
        st.font.name = FONT
        st.font.color.rgb = color
        st.font.italic = False
        rpr = st.element.get_or_add_rPr()
        rf = rpr.get_or_add_rFonts()
        for attr in list(rf.attrib):
            if attr.endswith("Theme"):
                del rf.attrib[attr]
        for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
            rf.set(qn(attr), FONT)
        color_el = rpr.find(qn("w:color"))
        if color_el is not None:
            for attr in list(color_el.attrib):
                if "theme" in attr.lower():
                    del color_el.attrib[attr]


def _styled_para(doc, text, size, color, bold=False, italic=False,
                 space_before=0, space_after=8, keep_with_next=False, align=None):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.keep_with_next = keep_with_next
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        r.font.name = FONT
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.bold = bold
        r.italic = italic
    return p


def title(doc, text, subtitle=None):
    _styled_para(doc, text, 24, NEAR_BLACK, bold=True, space_after=2, keep_with_next=True)
    if subtitle:
        _styled_para(doc, subtitle, 11, MID_GRAY, space_after=6, keep_with_next=True)
    accent_bar(doc)


def accent_bar(doc, width_in=6.5, space_after=14):
    """The purple rule. A one-cell table with a purple fill, per the brand skill -
    a Word border or an <hr> both come through as a grey hairline.

    The row height is pinned (exact, 2.5pt) and the paragraph mark is shrunk,
    because an empty paragraph takes its height from the mark, not the run: left
    at the 10pt default it drew a slab over 15pt tall instead of a rule. Pages
    floors every table row at 12pt, so the bar is 12pt there whatever the XML
    says; Apple's Word importer (Quick Look) draws the 2.5pt. The bar and the
    spacer after it keep with next, so a title never strands its bar at the foot
    of a page."""
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    row = t.rows[0]
    row.height = Twips(50)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    cell = row.cells[0]
    cell.width = Inches(width_in)
    _cell_shade(cell, HEX_PURPLE)
    _cell_margins(cell, 0, 0, 0, 0)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1
    p.paragraph_format.keep_with_next = True
    r = p.add_run("")
    r.font.size = Pt(2)
    mark = OxmlElement("w:rPr")
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), "2")
    mark.append(sz)
    p._p.get_or_add_pPr().insert_element_before(mark, "w:sectPr", "w:pPrChange")
    _grid_widths(t, [width_in])
    _spacer(doc, space_after, keep_with_next=True)
    return t


def _spacer(doc, pts, keep_with_next=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1
    if keep_with_next:
        p.paragraph_format.keep_with_next = True
    p.add_run("").font.size = Pt(pts / 2 if pts else 1)


def page_break(doc):
    from docx.enum.text import WD_BREAK
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_break(WD_BREAK.PAGE)


def _heading(doc, text, level):
    size, color, before, after = HEADING_LOOK[level]
    p = doc.add_paragraph(style=f"Heading {level}")
    pf = p.paragraph_format
    pf.space_before, pf.space_after, pf.keep_with_next = Pt(before), Pt(after), True
    r = p.add_run(text)
    r.font.name, r.font.size, r.font.color.rgb, r.bold, r.italic = FONT, Pt(size), color, True, False
    rf = r._element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(attr), FONT)
    return p


def h1(doc, text):
    """Section title (Heading 1) with the purple accent bar under it."""
    p = _heading(doc, text, 1)
    accent_bar(doc, space_after=10)
    return p


def h2(doc, text):
    return _heading(doc, text, 2)


def h3(doc, text):
    return _heading(doc, text, 3)


def _add_parts(p, parts, size=10):
    """Runs from a parts list. A part is a plain str, or a tuple:
    ("b", text) bold, ("pb", text) purple bold label, ("i", text) italic,
    ("c", text) a channel/skill/path name (bold, dark grey), ("link", text, url)."""
    if isinstance(parts, str):
        parts = [parts]
    for part in parts:
        if isinstance(part, str):
            kind, text = "", part
        else:
            kind, text = part[0], part[1]
        if kind == "link":
            add_hyperlink(p, part[2], text, size=size)
            continue
        r = p.add_run(text)
        r.font.name, r.font.size = FONT, Pt(size)
        r.font.color.rgb = PURPLE if kind == "pb" else DARK_GRAY if kind == "c" else NEAR_BLACK
        r.bold = kind in ("b", "pb", "c")
        r.italic = kind == "i"


def rich(doc, parts, size=10, space_after=9):
    """A body paragraph mixing bold, labels and real hyperlinks (see _add_parts)."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    _add_parts(p, parts, size=size)
    return p


def picture(doc, path, width_in=6.5, space_after=6):
    """A centred image at the given width (PNG/JPEG). Convert big screenshots to
    JPEG first: a 1920px PNG hero is ~700KB, the same shot as JPEG ~190KB."""
    doc.add_picture(path, width=Inches(width_in))
    p = doc.paragraphs[-1]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(space_after)
    return p


def body(doc, text, size=10, space_after=9, italic=False, color=NEAR_BLACK):
    return _styled_para(doc, text, size, color, italic=italic, space_after=space_after)


def caption(doc, text, space_after=9):
    return _styled_para(doc, text, 8.5, MID_GRAY, space_after=space_after)


def label_para(doc, label, text):
    """Purple bold label, then normal body in the same paragraph."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(9)
    r1 = p.add_run(label)
    r1.font.name, r1.font.size, r1.font.color.rgb, r1.bold = FONT, Pt(10), PURPLE, True
    r2 = p.add_run(text)
    r2.font.name, r2.font.size, r2.font.color.rgb = FONT, Pt(10), NEAR_BLACK
    return p


def _fresh_list_num(doc):
    """A numbering instance of its own for one numbered list, counting from 1.

    Every List Number paragraph inherits one numId from the style, and a numId is
    one counter, so a second numbered list carried on from the first: a 4-item
    question list after a 6-item list rendered as 7-10 in Pages and in the
    converted Google Doc (2026-09-29), which broke "see question 1".

    The list gets a new w:num with startOverride 1, which Pages honours, pointing
    at its own copy of the style's abstractNum, because Apple's Word importer
    (Quick Look) ignores startOverride and counts every list that shares an
    abstractNum as one. The copy gets a fresh nsid and drops its level's pStyle,
    so the style stays linked to one list only. Returns the new numId."""
    numbering = doc.part.numbering_part.element
    style_num = numbering.num_having_numId(doc.styles["List Number"].element.pPr.numPr.numId.val)
    abstracts = numbering.findall(qn("w:abstractNum"))
    src = next(a for a in abstracts if a.get(qn("w:abstractNumId")) == str(style_num.abstractNumId.val))
    abs_id = max(int(a.get(qn("w:abstractNumId"))) for a in abstracts) + 1
    clone = copy.deepcopy(src)
    clone.set(qn("w:abstractNumId"), str(abs_id))
    nsid = clone.find(qn("w:nsid"))
    if nsid is not None:
        nsid.set(qn("w:val"), f"{0x52560000 + abs_id:08X}")
    for ps in clone.findall(f"{qn('w:lvl')}/{qn('w:pStyle')}"):
        ps.getparent().remove(ps)
    abstracts[-1].addnext(clone)  # every abstractNum precedes every w:num
    num = numbering.add_num(abs_id)
    num.add_lvlOverride(ilvl=0).add_startOverride(1)
    return num.numId


def _list_num_id(doc, restart):
    """numId for the next numbered list: a fresh counter, or with restart=False the
    one the last numbered paragraph used, so the count carries on."""
    if not restart:
        sid = doc.styles["List Number"].style_id
        prev = doc.element.body.xpath(
            f'./w:p[w:pPr/w:pStyle/@w:val="{sid}"]/w:pPr/w:numPr/w:numId/@w:val')
        if prev:
            return int(prev[-1])
    return _fresh_list_num(doc)


def bullets(doc, items, numbered=False, space_after=4, keep_together=False, restart=True):
    """Each item is a plain str or a parts list (see _add_parts), so a bullet can
    carry a bold label or a working link. keep_together keeps a short list on one
    page, so its last item does not strand overleaf. A numbered list starts at 1;
    restart=False continues the previous numbered list's count instead."""
    style = "List Number" if numbered else "List Bullet"
    num_id = _list_num_id(doc, restart) if numbered else None
    for i, item in enumerate(items):
        p = doc.add_paragraph(style=style)
        if num_id is not None:  # schema-ordered insert: after pStyle, before spacing
            numPr = p._p.get_or_add_pPr().get_or_add_numPr()
            numPr.get_or_add_ilvl().val = 0
            numPr.get_or_add_numId().val = num_id
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        if keep_together and i < len(items) - 1:
            p.paragraph_format.keep_with_next = True
        _add_parts(p, item)
    _spacer(doc, 8)


def callout(doc, text, size=13, fill=HEX_PURPLE_MIST):
    """Goal box: purple left bar, mist fill, real padding. The HTML equivalent
    flattens to highlighted text, which is what made the first one-pager look
    like someone had taken a marker to it."""
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    cell = t.rows[0].cells[0]
    cell.width = Inches(6.5)
    _cell_shade(cell, fill)
    _cell_margins(cell, 180, 180, 220, 220)
    _left_bar(cell)
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    r.font.name, r.font.size, r.font.color.rgb = FONT, Pt(size), NEAR_BLACK
    _row_keep_together(t.rows[0])
    _grid_widths(t, [6.5])
    _spacer(doc, 16)
    return t


def kv_table(doc, pairs, key_width=1.6, total=6.5):
    t = doc.add_table(rows=0, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    _table_borders(t)
    for k, v in pairs:
        row = t.add_row()
        _row_keep_together(row)
        kc, vc = row.cells
        kc.width, vc.width = Inches(key_width), Inches(total - key_width)
        _cell_shade(kc, HEX_PURPLE_LIGHT)
        for cell, text, bold in ((kc, k, False), (vc, v, False)):
            _cell_margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            r = p.add_run(text)
            r.font.name, r.font.size, r.font.color.rgb = FONT, Pt(9.5), NEAR_BLACK
            r.bold = bold
    _grid_widths(t, [key_width, total - key_width])
    _spacer(doc, 10)
    return t


def table(doc, headers, rows, widths=None, size=9, check_col=False):
    ncols = len(headers)
    widths = widths or [6.5 / ncols] * ncols
    t = doc.add_table(rows=0, cols=ncols)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    _table_borders(t)

    hrow = t.add_row()
    _row_keep_together(hrow, header=True)
    for cell, text, w in zip(hrow.cells, headers, widths):
        cell.width = Inches(w)
        _cell_shade(cell, HEX_NEAR_BLACK)
        _cell_margins(cell)
        p = cell.paragraphs[0]
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(text)
        r.font.name, r.font.size, r.font.color.rgb, r.bold = FONT, Pt(size), WHITE, True

    for i, data in enumerate(rows):
        row = t.add_row()
        _row_keep_together(row)
        for col, (cell, text, w) in enumerate(zip(row.cells, data, widths)):
            cell.width = Inches(w)
            if i % 2 == 1:
                _cell_shade(cell, HEX_PURPLE_LIGHT)
            _cell_margins(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.1
            # NB: `cell is row.cells[0]` does NOT work - row.cells builds a new
            # tuple of new _Cell objects on every access, so identity never matches.
            is_check = check_col and col == 0
            if is_check:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(text)
            r.font.name = FONT
            r.font.size = Pt(13 if is_check else size)
            r.font.color.rgb = MID_GRAY if is_check else NEAR_BLACK
    _grid_widths(t, widths)
    _spacer(doc, 12)
    return t


def answer_box(doc, lines=3):
    """An empty bordered box for someone to type into."""
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    cell = t.rows[0].cells[0]
    cell.width = Inches(6.5)
    _cell_margins(cell, 120, 120, 140, 140)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.add_run("").font.size = Pt(10)
    for _ in range(lines - 1):
        q = cell.add_paragraph()
        q.paragraph_format.space_after = Pt(0)
        q.add_run("").font.size = Pt(10)
    _table_borders(t)
    _grid_widths(t, [6.5])
    _spacer(doc, 12)
    return t


def link_line(doc, parts, size=8.5, space_before=4):
    """A paragraph mixing plain text and real hyperlinks.
    parts: list of str, or (text, url) tuples."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(8)
    for part in parts:
        if isinstance(part, tuple):
            add_hyperlink(p, part[1], part[0], size=size)
        else:
            r = p.add_run(part)
            r.font.name, r.font.size, r.font.color.rgb = FONT, Pt(size), MID_GRAY
    return p


DILIGENCE = (
    "In creating this {kind}, I collaborated with Claude, an AI assistant by Anthropic, "
    "to assist with {tasks}. I affirm that all AI-generated and co-created content "
    "underwent thorough review and evaluation. The final output accurately reflects my "
    "understanding, expertise, and intended meaning. While AI assistance was instrumental "
    "in the process, I maintain full responsibility for the content, its accuracy, and its "
    "presentation. This disclosure is made in the spirit of transparency and to acknowledge "
    "the role of AI in the creation process."
)


def diligence(doc, kind="document", tasks="drafting and formatting"):
    _spacer(doc, 18)
    accent_bar(doc, space_after=8)
    _styled_para(doc, "AI diligence statement", 10, PURPLE, bold=True,
                 space_after=3, keep_with_next=True)
    _styled_para(doc, DILIGENCE.format(kind=kind, tasks=tasks), 8, NEAR_BLACK, space_after=0)


def brand_header(doc, right_text=""):
    section = doc.sections[0]
    p = section.header.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r1 = p.add_run("Riverside.com")
    r1.font.name, r1.font.size, r1.font.color.rgb, r1.bold = FONT, Pt(8.5), PURPLE, True
    if right_text:
        r2 = p.add_run("   " + right_text)
        r2.font.name, r2.font.size, r2.font.color.rgb = FONT, Pt(8.5), MID_GRAY


def page_numbers(doc):
    section = doc.sections[0]
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = p.add_run()
    run.font.name, run.font.size, run.font.color.rgb = FONT, Pt(8), MID_GRAY
    for instr in ('begin', 'PAGE', 'end'):
        el = OxmlElement("w:fldChar" if instr != "PAGE" else "w:instrText")
        if instr == "PAGE":
            el.set(qn("xml:space"), "preserve")
            el.text = " PAGE "
        else:
            el.set(qn("w:fldCharType"), instr)
        run._r.append(el)


# ------------------------------------------------------------------ packaging

def slim(src_path, dst_path):
    """Strip python-docx default-template bloat so the file fits an inline upload.

    stylesWithEffects.xml is a legacy Word 2010 part nothing reads, and styles.xml
    ships ~370 style definitions of which a generated document uses a handful.
    Dropping the rest is safe as long as every style id referenced by document.xml,
    numbering.xml and the headers survives. Returns (bytes_before, bytes_after).
    numbering.xml itself is copied as is: it holds the w:num each numbered list
    restarts from (see _fresh_list_num).

    Both the content types map and BOTH rels files must be scrubbed; the thumbnail
    is referenced from _rels/.rels, not word/_rels/document.xml.rels, and missing
    that one produces a file python-docx itself cannot reopen.
    """
    import os, re, shutil, tempfile, zipfile
    from lxml import etree

    W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    DROP = {"word/stylesWithEffects.xml", "docProps/thumbnail.jpeg"}
    KEEP_STYLES = {"Normal", "ListParagraph", "ListBullet", "ListNumber", "TableGrid",
                   "Header", "Footer", "Hyperlink", "DefaultParagraphFont",
                   "NoList", "TableNormal"}

    before = os.path.getsize(src_path)
    zin = zipfile.ZipFile(src_path)
    used = set()
    for n in zin.namelist():
        if n.endswith(".xml") and any(k in n for k in ("document", "numbering", "header", "footer")):
            used.update(m.group(1).decode() for m in re.finditer(rb'w:val="([^"]+)"', zin.read(n)))

    root = etree.fromstring(zin.read("word/styles.xml"))
    # A kept paragraph style can name a linked character style (Heading 1 ->
    # Heading1Char); keep that too, so no w:link is left pointing at nothing.
    for st in root.findall(f"{W}style"):
        link = st.find(f"{W}link")
        if st.get(f"{W}styleId") in used and link is not None:
            used.add(link.get(f"{W}val"))
    for st in list(root.findall(f"{W}style")):
        sid = st.get(f"{W}styleId")
        if not (sid in used or st.get(f"{W}default") == "1" or sid in KEEP_STYLES):
            root.remove(st)
    new_styles = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)

    tmp = tempfile.mkdtemp()
    try:
        for item in zin.infolist():
            if item.filename in DROP:
                continue
            target = os.path.join(tmp, item.filename)
            os.makedirs(os.path.dirname(target), exist_ok=True)
            data = new_styles if item.filename == "word/styles.xml" else zin.read(item.filename)
            with open(target, "wb") as fh:
                fh.write(data)
        zin.close()

        ct = os.path.join(tmp, "[Content_Types].xml")
        s = open(ct, encoding="utf-8").read()
        for pat in (r'<Override[^>]*stylesWithEffects[^>]*/>',
                    r'<Override[^>]*thumbnail[^>]*/>',
                    r'<Default[^>]*Extension="jpeg"[^>]*/>'):
            s = re.sub(pat, "", s)
        open(ct, "w", encoding="utf-8").write(s)

        for rels in (os.path.join(tmp, "word", "_rels", "document.xml.rels"),
                     os.path.join(tmp, "_rels", ".rels")):
            if not os.path.exists(rels):
                continue
            r = open(rels, encoding="utf-8").read()
            for pat in (r'<Relationship[^>]*stylesWithEffects[^>]*/>',
                        r'<Relationship[^>]*thumbnail[^>]*/>'):
                r = re.sub(pat, "", r)
            open(rels, "w", encoding="utf-8").write(r)

        zf = zipfile.ZipFile(dst_path, "w", zipfile.ZIP_DEFLATED, compresslevel=9)
        for base, _, files in os.walk(tmp):
            for f in files:
                full = os.path.join(base, f)
                zf.write(full, os.path.relpath(full, tmp))
        zf.close()
    finally:
        shutil.rmtree(tmp)
    return before, os.path.getsize(dst_path)


def _self_test():
    """Build a document exercising every helper, slim it, and reopen it."""
    import os, tempfile
    from docx import Document as _D
    import base64
    d = new_doc()
    brand_header(d, "Self test")
    page_numbers(d)
    title(d, "Self test", "every helper once")
    callout(d, "A callout box.")
    kv_table(d, [("Key", "Value")])
    h1(d, "A top section"); h2(d, "A section"); h3(d, "A subsection")
    body(d, "Body text."); caption(d, "A caption.")
    label_para(d, "Label: ", "and the rest.")
    rich(d, [("pb", "Label: "), "text with ", ("b", "bold"), " and a ", ("link", "link", "https://riverside.com")])
    bullets(d, ["one", "two"]); bullets(d, ["first", "second"], numbered=True)
    bullets(d, ["first again", "second again"], numbered=True)  # restarts at 1
    bullets(d, ["third again"], numbered=True, restart=False)   # carries on at 3
    bullets(d, [[("c", "#channel"), " and ", ("link", "a link", "https://riverside.com/pricing")], "plain"],
            keep_together=True)
    png = os.path.join(tempfile.mkdtemp(), "px.png")  # 1x1 PNG, no imaging library needed
    with open(png, "wb") as fh:
        fh.write(base64.b64decode("iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg=="))
    picture(d, png, width_in=1)
    table(d, ["A", "B"], [["1", "2"], ["3", "4"]])
    check_widths = [0.5, 6.0]
    table(d, ["Need", "Item"], [["☐", "x"]], widths=check_widths, check_col=True)
    answer_box(d, 2)
    page_break(d)
    link_line(d, ["See ", ("Riverside", "https://riverside.com"), "."])
    diligence(d, "document", "drafting")
    tmp = tempfile.mkdtemp()
    raw, small = os.path.join(tmp, "t.docx"), os.path.join(tmp, "t.slim.docx")
    d.save(raw)
    before, after = slim(raw, small)

    checks = []
    r = _D(small)
    checks.append(("reopens after slim", True))
    checks.append(("tables present", len(r.tables) >= 6))
    checks.append(("shrank", after < before))
    purple = sum(1 for p in r.paragraphs for run in p.runs
                 if run.font.color and run.font.color.rgb and str(run.font.color.rgb) == HEX_PURPLE)
    checks.append(("purple headings survive", purple >= 2))
    fills = r.element.xml
    checks.append(("purple accent fill present", HEX_PURPLE in fills))
    checks.append(("alt-row tint present", HEX_PURPLE_LIGHT in fills))
    checks.append(("header-row repeat set", "tblHeader" in fills))
    checks.append(("rows kept whole", "cantSplit" in fills))
    checks.append(("hyperlink relationship", any("hyperlink" in rel.reltype for rel in r.part.rels.values())))
    big = [run for t in r.tables for row in t.rows for run in row.cells[0].paragraphs[0].runs
           if run.text.strip() == "☐" and run.font.size and run.font.size >= Pt(12)]
    checks.append(("check glyph legible", len(big) >= 1))

    def tags(el):
        return [c.tag.split("}")[1] for c in el]

    def in_order(names):  # known tags, strictly increasing: no strays, no duplicates
        if not all(n in TBLPR_ORDER for n in names):
            return False
        idx = [TBLPR_ORDER.index(n) for n in names]
        return all(a < b for a, b in zip(idx, idx[1:]))

    checks.append(("tblPr children in schema order",
                   all(in_order(tags(t._tbl.tblPr)) for t in r.tables)))
    need = next(t for t in r.tables if t.cell(0, 0).text == "Need")
    grid = [int(g.get(qn("w:w"))) for g in need._tbl.tblGrid.findall(qn("w:gridCol"))]
    tblW = need._tbl.tblPr.find(qn("w:tblW"))
    want = [int(round(w * 1440)) for w in check_widths]
    checks.append(("gridCol widths match requested", grid == want
                   and tblW.get(qn("w:type")) == "dxa" and int(tblW.get(qn("w:w"))) == sum(want)))
    bars = [t for t in r.tables if len(t.columns) == 1 and f'w:fill="{HEX_PURPLE}"' in t._tbl.xml]
    checks.append(("accent bar row height exact", len(bars) >= 2 and all(
        b.rows[0].height_rule == WD_ROW_HEIGHT_RULE.EXACTLY and b.rows[0].height == Twips(50)
        for b in bars)))
    checks.append(("accent bar keeps with next", len(bars) >= 2 and all(
        b.cell(0, 0).paragraphs[0].paragraph_format.keep_with_next for b in bars)))

    used_styles = {p.style.name for p in r.paragraphs}
    checks.append(("h1-h3 use real Heading styles (document outline)",
                   {"Heading 1", "Heading 2", "Heading 3"} <= used_styles))
    style_ids = {s.style_id for s in r.styles}
    theme_free, links_ok = True, True
    for level in (1, 2, 3):
        el = r.styles[f"Heading {level}"].element
        rf = el.find(qn("w:rPr") + "/" + qn("w:rFonts"))
        if rf is None or any(a.endswith("Theme") for a in rf.attrib):
            theme_free = False
        link = el.find(qn("w:link"))
        if link is not None and link.get(qn("w:val")) not in style_ids:
            links_ok = False
    checks.append(("heading styles carry no theme fonts", theme_free))
    checks.append(("heading linked styles survive slim", links_ok))
    checks.append(("picture embedded", len(r.inline_shapes) >= 1))
    checks.append(("link inside a bullet", any(
        p.style.name == "List Bullet" and p._p.find(qn("w:hyperlink")) is not None for p in r.paragraphs)))

    numbering = r.part.numbering_part.element
    style_num = r.styles["List Number"].element.pPr.numPr.numId.val
    style_abs = numbering.xpath(f'./w:num[@w:numId="{style_num}"]/w:abstractNumId/@w:val')

    def list_ids(*texts):  # the numId on each paragraph of one list, None if it has none
        return [(p._p.xpath("./w:pPr/w:numPr/w:numId/@w:val") or [None])[0]
                for p in r.paragraphs if p.text in texts]

    def abstract_of(num_id):
        return numbering.xpath(f'./w:num[@w:numId="{num_id}"]/w:abstractNumId/@w:val')

    def restarts(num_id):  # own w:num with startOverride 1, own decimal abstract list
        lvl = f'./w:abstractNum[@w:abstractNumId="{(abstract_of(num_id) or [""])[0]}"]/w:lvl[@w:ilvl="0"]'
        return (abstract_of(num_id) not in ([], style_abs)
                and numbering.xpath(lvl + "/w:numFmt/@w:val") == ["decimal"]
                and numbering.xpath(f'./w:num[@w:numId="{num_id}"]/w:lvlOverride[@w:ilvl="0"]'
                                    "/w:startOverride/@w:val") == ["1"])

    first, again = list_ids("first", "second"), list_ids("first again", "second again")
    checks.append(("each numbered list restarts at 1", len(first) == len(again) == 2
                   and len(set(first)) == len(set(again)) == 1 and None not in first + again
                   and first[0] != again[0] and abstract_of(first[0]) != abstract_of(again[0])
                   and restarts(first[0]) and restarts(again[0])))
    nsids = numbering.xpath("./w:abstractNum/w:nsid/@w:val")
    order = "".join({"abstractNum": "a", "num": "n"}.get(t, "") for t in tags(numbering))
    checks.append(("list copies: unique nsids, one style link, abstracts before nums",
                   len(nsids) == len(set(nsids)) and order == "a" * order.count("a") + "n" * order.count("n")
                   and len(numbering.xpath('./w:abstractNum[w:lvl/w:pStyle/@w:val="ListNumber"]')) == 1))
    checks.append(("restart=False continues the previous list",
                   None not in again and list_ids("third again") == again[:1]))
    numbered = [tags(p._p.pPr) for p in r.paragraphs if p.style.name == "List Number"]
    checks.append(("numPr after pStyle, before spacing", len(numbered) == 5 and all(
        t[0] == "pStyle" and {"numPr", "spacing"} <= set(t) and t.index("numPr") < t.index("spacing")
        for t in numbered)))

    failed = [n for n, ok in checks if not ok]
    for n, ok in checks:
        print(f"  {'ok   ' if ok else 'FAIL '} {n}")
    print(f"\n  size {before} -> {after} bytes")
    return 1 if failed else 0


if __name__ == "__main__":
    import sys
    sys.exit(_self_test())
